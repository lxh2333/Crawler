import codecs
import re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from bs4 import BeautifulSoup
from scrapy import Spider, Request
from selenium import webdriver
from pyquery import PyQuery as pq
import time


class HuaweiSpider(Spider):
    name = 'Huawei_news_en'
    allowed_domains = ['www.huawei.com']

    # custom_settings = {
    # 'LOG_LEVEL': 'INFO',
    # }

    def __init__(self):
        SERVICE_ARGS = ['--load-images=false', '--disk-cache=true', '--ignore-ssl-errors=true']
        self.broswer = webdriver.PhantomJS(service_args=SERVICE_ARGS,
                                           executable_path=r'C:\phantomjs-2.1.1-windows\phantomjs-2.1.1-windows\bin\phantomjs.exe')
        self.broswer.set_page_load_timeout(30)

    def closed(self, spider):
        self.broswer.close()

    def start_requests(self):
        url_css = 'https://www.huawei.com/en/press-events/news?d=ws&pagesize=100&pageindex={}'
        for i in range(1, 23):
            yield Request(url=url_css.format(i), callback=self.parse)

    def parse(self, response):
        cons = []
        new_cons = []
        flag = []
        html = response.text
        content = re.sub(r'\n|\r|\t', '', html)
        re_script = re.compile(r'<script.*?>.*?</script>', re.I)
        content = re_script.sub('', content)
        re_style = re.compile(r'<style.*?>.*?</style>', re.I)
        content = re_style.sub('', content)
        re_nav = re.compile(r'<li>.*?<span>.*?</span>.*?<a link.*?</a></li>', re.I)
        content = re_nav.sub('', content)
        re_small = re.compile(r'<small>.*?</small>', re.I)
        content = re_small.sub('', content)
        re_policy = re.compile(r'<.*?class="browsehappy ReadPolicy">(.*?)</div>', re.I)
        content = re_policy.sub('', content)
        re_nav = re.compile(r'<div class="nav-gblnav hidden-xs  hidden-sm">(.*?)</div>', re.I)
        content = re_nav.sub('', content)
        re_affix = re.compile(r'<header class="affix-top">(.*?)</header>', re.I)
        content = re_affix.sub('', content)
        re_footer = re.compile(r'<footer>(.*?)</footer>', re.I)
        content = re_footer.sub('', content)
        re_con = re.compile(r'<div class="container">(.*?)</div>', re.I)
        content = re_con.sub('', content)
        re_global = re.compile(r'<div class="worldwide-language">(.*?)</div>', re.I)
        content = re_global.sub('', content)
        re_blank_next = re.compile(r'\s{2}', re.I)
        content = re_blank_next.sub('bupt', content)
        soup = BeautifulSoup(content, 'lxml')
        contents = soup.text
        sentences = contents.split('bupt')
        start = -1
        end = len(sentences)
        for sentence in sentences:
            start += 1
            if re.match(r'NEWS |.*?,.*？', sentence):
                break
        for i in range(end-start):
            if sentences[start+i].strip() == '' or sentences[start+i].strip() == 'Learn More' or sentences[start+i].strip() == 'Visit Event Website (External)':
                continue
            cons.append(sentences[start+i].strip())
        for i in range(len(cons)):
            if re.match(r'NEWS |.*?,.*？', cons[i]):
                flag.append(i)
        for i in range(len(flag)):
            if i != len(flag)-1:
                l = flag[i+1] - flag[i]
            if i == len(flag)-1:
                l = len(cons) - flag[i]
            cont = ''
            for j in range(l):
                if j != l-1:
                    cont += cons[flag[i]+j]+'|'
                else:
                    cont += cons[flag[i]+j]
            new_cons.append(cont)
        for new_con in new_cons:
            print(new_con)
        f = codecs.open('news_en.txt', 'a', 'utf-8')
        for new_con in new_cons:
            f.write(new_con + '\n')
        f.close()
        '''
        document = Document()
        document.styles['Normal'].font.name = u'宋体'
        document.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
        rows = 1
        cols = 2
        table = document.add_table(rows=1, cols=2, style="Table Grid")
        for i in range(rows):
            tr = table.rows[i]._tr
            trPr = tr.get_or_add_trPr()
            trHeight = OxmlElement('w:trHeight')
            trHeight.set(qn('w:val'), "450")
            trPr.append(trHeight)
        table.cell(0, 0).text = u'中文'
        table.cell(0, 1).text = u'英文'
        now = 1
        for sentence in sentences:
            if not sentence.strip() == '':
                sentence = sentence.strip(' ')
                if sentence != '|' and sentence != '×':
                    table.add_row()
                    table.cell(now, 0).text = sentence
                    now = now + 1
        document.save('test.docx')
        '''