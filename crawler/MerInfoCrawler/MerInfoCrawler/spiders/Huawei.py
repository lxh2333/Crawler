import codecs
import re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from bs4 import BeautifulSoup
from scrapy import Spider,Request
from selenium import webdriver
from pyquery import PyQuery as pq
import time
class HuaweiSpider(Spider):
    name = 'Huawei'
    allowed_domains = ['www.huawei.com']
    #custom_settings = {
        #'LOG_LEVEL': 'INFO',
    #}

    def __init__(self):
        SERVICE_ARGS = ['--load-images=false', '--disk-cache=true', '--ignore-ssl-errors=true']
        self.broswer=webdriver.PhantomJS(service_args=SERVICE_ARGS, executable_path = r'C:\phantomjs-2.1.1-windows\phantomjs-2.1.1-windows\bin\phantomjs.exe')
        self.broswer.set_page_load_timeout(30)

    def closed(self,spider):
        self.broswer.close()

    def start_requests(self):
        url_css =  'http://carrier.huawei.com/en/events/mwcs2018'
        yield Request(url=url_css, callback=self.parse)

    def parse(self, response):
        new_contents = []
        html = self.broswer.page_source
        html1 = response.text
        content = re.sub(r'\n|\r|\t', '', html1)
        re_http_raw = re.compile(r'<a href="(https://.*?)"')
        content_https = re_http_raw.findall(content)
        content_https = list(set(content_https))
        '''
        for content_http in content_https:
            if not content_http == '':
                f = codecs.open('https.txt', 'a', 'utf-8')
                f.write(content_http+'\n')
                f.close()
                '''
        re_script = re.compile(r'<script.*?>.*?</script>', re.I)
        content = re_script.sub('', content)
        re_style = re.compile(r'<style.*?>.*?</style>', re.I)
        content = re_style.sub('', content)
        re_nav = re.compile(r'<li>.*?<span>.*?</span>.*?<a link.*?</a></li>', re.I)
        content = re_nav.sub('', content)
        re_foot = re.compile(r'<div class="row">(.*?)</div>', re.I)
        content = re_foot.sub('', content)
        re_small = re.compile(r'<small>.*?</small>', re.I)
        content = re_small.sub('', content)
        re_policy = re.compile(r'<.*?class="browsehappy ReadPolicy">(.*?)</div>', re.I)
        content = re_policy.sub('', content)
        re_nav = re.compile(r'<div class="nav-gblnav hidden-xs  hidden-sm">(.*?)</div>', re.I)
        content = re_nav.sub('', content)
        re_affix = re.compile(r'<header class="affix-top">(.*?)</header>', re.I)
        content = re_affix.sub('', content)
        re_footer = re.compile(r'<footer>(.*?)</footer>', re.I)
        content = re_footer.sub('', content)
        re_con = re.compile(r'<div class="container">(.*?)</div>', re.I)
        content = re_con.sub('', content)
        re_global = re.compile(r'<div class="worldwide-language">(.*?)</div>', re.I)
        content = re_global.sub('', content)
        re_blank_next = re.compile(r'\s{3}', re.I)
        content = re_blank_next.sub('bupt', content)
        soup = BeautifulSoup(content, 'lxml')
        contents = soup.text
        sentences = contents.split('bupt')
        for sentence in sentences:
            if not sentence.strip() == '':
                sentence = sentence.strip(' ')
                f = codecs.open('en_test.txt', 'a', 'utf-8')
                if sentence != '|' and sentence != '×':
                    f.write(sentence + '\n')
                    f.close()
        '''
        for sentence in sentences:
            if not sentence.strip() == '':
                sentence = sentence.strip(' ')
                f = codecs.open('cn_1.txt', 'a', 'utf-8')
                if sentence != '|' and sentence != '×':
                    f.write(sentence+'\n')
                    f.close()
        
        document = Document()
        document.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
        rows = 1
        cols = 2
        table = document.add_table(rows=1, cols=2, style="Table Grid")
        for i in range(rows):
            tr = table.rows[i]._tr
            trPr = tr.get_or_add_trPr()
            trHeight = OxmlElement('w:trHeight')
            trHeight.set(qn('w:val'), "450")
            trPr.append(trHeight)
        table.cell(0, 0).text = u'中文'
        table.cell(0, 1).text = u'英文'
        now = 1
        for sentence in sentences:
            if not sentence.strip() == '':
                sentence = sentence.strip(' ')
                if sentence != '|' and sentence != '×':
                    table.add_row()
                    table.cell(now, 0).text = sentence
                    now = now + 1
        document.save('test.docx')
        '''