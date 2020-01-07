import codecs
import os
import re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from bs4 import BeautifulSoup
from scrapy import Spider, Request
from selenium import webdriver
from ChineseSimilartyCaculation import CSV
from translate_google import google_translate




class HuaweiNextSpider(Spider):
    name = 'Huawei_next'
    allowed_domains = []
    cn_contents = []
    en_contents = []
    # custom_settings = {
    # 'LOG_LEVEL': 'INFO',
    # }

    def __init__(self, url=None):
        SERVICE_ARGS = ['--load-images=false', '--disk-cache=true', '--ignore-ssl-errors=true']
        self.broswer = webdriver.PhantomJS(service_args=SERVICE_ARGS,
                                           executable_path=r'C:\phantomjs-2.1.1-windows\phantomjs-2.1.1-windows\bin\phantomjs.exe')
        self.broswer.set_page_load_timeout(30)
        self.url = url

    def closed(self, spider):
        self.broswer.close()

    def start_requests(self):
        '''
        urls = []
        for url in open('href_cases.txt'):
            if re.match(r'.*?/.*?cn/.*?', url):
                self.cn_contents = []
                self.en_contents = []
                yield Request(url=url, callback=self.parse)
                print(url)
                url = url.replace("cn", "en")
                yield Request(url=url, callback=self.en_parse)
                print(url)
        '''
        yield Request(url=self.url, callback=self.parse)
        url = self.url.replace("cn", "en")
        yield Request(url=url, callback=self.en_parse)

    def Num_en(self, sentences):
        sentences = sentences.strip()
        index = 0
        count = 0
        while index < len(sentences):
            while sentences[index] != " ":
                index += 1
                if index == len(sentences):
                    break
            count += 1
            if index == len(sentences):
                break
            while sentences[index] == " ":
                index += 1
        return count

    def Doc(self, path_cn, content_en, doc):
        en = []
        cn = []
        cn_pick = []


        flag = -1
        for sentence in content_en:
            if not sentence.strip() == '':
                sentence = sentence.strip(' ')
                if sentence != '|' and sentence != '×':
                    flag = flag + 1
                    if self.Num_en(sentence) > 10:
                        trans = google_translate(sentence)
                        #csv = CSV(trans, 'news.txt')
                        # print("最相似的文本为：",csv)
                        csv = CSV(trans, path_cn)
                        if csv != 0:
                            en.append(sentence)
                            cn_pick.append(csv)
        for i in range(len(en)):
            print(en[i])
            print("最相似的文本为：", cn_pick[i])
        '''
        document = Document(doc)
        table = document.tables[0]
        rows = len(table.rows)
        for i in range(len(en)):
            table.add_row()
            table.cell(rows+i, 0).text = cn_pick[i]
            table.cell(rows+i, 1).text = en[i]
        document.save(doc)
        '''

    def en_parse(self, response):
        html1 = response.text
        content = re.sub(r'\n|\r|\t', '', html1)
        re_http_raw = re.compile(r'<a href="(https://.*?)"')
        content_https = re_http_raw.findall(content)
        content_https = list(set(content_https))
        re_script = re.compile(r'<script.*?>.*?</script>', re.I)
        content = re_script.sub('', content)
        re_style = re.compile(r'<style.*?>.*?</style>', re.I)
        content = re_style.sub('', content)
        re_nav = re.compile(r'<li>.*?<span>.*?</span>.*?<a link.*?</a></li>', re.I)
        content = re_nav.sub('', content)
        re_foot = re.compile(r'<div class="row">(.*?)</div>', re.I)
        content = re_foot.sub('', content)
        re_small = re.compile(r'<small>.*?</small>', re.I)
        content = re_small.sub('', content)
        re_policy = re.compile(r'<.*?class="browsehappy ReadPolicy">(.*?)</div>', re.I)
        content = re_policy.sub('', content)
        re_nav = re.compile(r'<div class="nav-gblnav hidden-xs  hidden-sm">(.*?)</div>', re.I)
        content = re_nav.sub('', content)
        re_affix = re.compile(r'<header class="affix-top">(.*?)</header>', re.I)
        content = re_affix.sub('', content)
        re_footer = re.compile(r'<footer>(.*?)</footer>', re.I)
        content = re_footer.sub('', content)
        re_con = re.compile(r'<div class="container">(.*?)</div>', re.I)
        content = re_con.sub('', content)
        re_global = re.compile(r'<div class="worldwide-language">(.*?)</div>', re.I)
        content = re_global.sub('', content)
        re_blank_next = re.compile(r'\s{3}', re.I)
        content = re_blank_next.sub('bupt', content)
        soup = BeautifulSoup(content, 'lxml')
        contents = soup.text
        sentences = contents.split('bupt')
        IsExt = os.path.isfile('cn.txt')
        if IsExt:
            os.remove('en.txt')
        f = codecs.open('en.txt', 'a', 'utf-8')
        for sentence in sentences:
            if not sentence.strip() == '':
                sentence = sentence.strip(' ')
                if sentence != '|' and sentence != '×':
                    self.en_contents.append(sentence)
                    f.write(sentence + '\n')

    def parse(self, response):
        html1 = response.text
        content = re.sub(r'\n|\r|\t', '', html1)
        re_http_raw = re.compile(r'<a href="(https://.*?)"')
        content_https = re_http_raw.findall(content)
        content_https = list(set(content_https))
        '''
        for content_http in content_https:
            if not content_http == '':
                f = codecs.open('https.txt', 'a', 'utf-8')
                f.write(content_http+'\n')
                f.close()
        '''
        re_script = re.compile(r'<script.*?>.*?</script>', re.I)
        content = re_script.sub('', content)
        re_style = re.compile(r'<style.*?>.*?</style>', re.I)
        content = re_style.sub('', content)
        re_nav = re.compile(r'<li>.*?<span>.*?</span>.*?<a link.*?</a></li>', re.I)
        content = re_nav.sub('', content)
        re_foot = re.compile(r'<div class="row">(.*?)</div>', re.I)
        content = re_foot.sub('', content)
        re_small = re.compile(r'<small>.*?</small>', re.I)
        content = re_small.sub('', content)
        re_policy = re.compile(r'<.*?class="browsehappy ReadPolicy">(.*?)</div>', re.I)
        content = re_policy.sub('', content)
        re_nav = re.compile(r'<div class="nav-gblnav hidden-xs  hidden-sm">(.*?)</div>', re.I)
        content = re_nav.sub('', content)
        re_affix = re.compile(r'<header class="affix-top">(.*?)</header>', re.I)
        content = re_affix.sub('', content)
        re_footer = re.compile(r'<footer>(.*?)</footer>', re.I)
        content = re_footer.sub('', content)
        re_con = re.compile(r'<div class="container">(.*?)</div>', re.I)
        content = re_con.sub('', content)
        re_global = re.compile(r'<div class="worldwide-language">(.*?)</div>', re.I)
        content = re_global.sub('', content)
        re_blank_next = re.compile(r'\s{3}', re.I)
        content = re_blank_next.sub('bupt', content)
        soup = BeautifulSoup(content, 'lxml')
        contents = soup.text
        sentences = contents.split('bupt')
        IsExt = os.path.isfile('cn.txt')
        if IsExt:
            os.remove('cn.txt')
        f = codecs.open('cn.txt', 'a', 'utf-8')
        for sentence in sentences:
            if not sentence.strip() == '':
                sentence = sentence.strip(' ')
                if sentence != '|' and sentence != '×':
                    f.write(sentence + '\n')
        f.close()