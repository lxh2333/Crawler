import codecs
import re

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from bs4 import BeautifulSoup
from scrapy import Spider,Request
from selenium import webdriver
from pyquery import PyQuery as pq
import time
import urllib.request
import urllib.parse
import json
class HuaweiENSpider(Spider):
    name = 'Huawei_EN'
    allowed_domains = ['www.huawei.com']
    #custom_settings = {
        #'LOG_LEVEL': 'INFO',
    #}

    def __init__(self):
        SERVICE_ARGS = ['--load-images=false','--disk-cache=true','--ignore-ssl-errors=true']
        self.broswer=webdriver.PhantomJS(service_args=SERVICE_ARGS,executable_path=r'C:\phantomjs-2.1.1-windows\phantomjs-2.1.1-windows\bin\phantomjs.exe')
        self.broswer.set_page_load_timeout(30)
    def closed(self,spider):
        self.broswer.close()
    def start_requests(self):
        url_css =  'https://www.huawei.com/en/'
        yield Request(url=url_css,callback=self.parse)
    def parse(self, response):
        html = self.broswer.page_source
        html1 = response.text
        content = re.sub(r'\n|\r|\t', '', html1)
        re_script = re.compile(r'<script.*?>.*?</script>', re.I)
        content = re_script.sub('', content)
        re_style = re.compile(r'<style.*?>.*?</style>', re.I)
        content = re_style.sub('', content)
        re_nav = re.compile(r'<li>.*?<span>.*?</span>.*?<a link.*?</a></li>', re.I)
        content = re_nav.sub('', content)
        re_foot = re.compile(r'<div class="row">.*?</div>', re.I)
        content = re_foot.sub('', content)
        re_policy = re.compile(r'<.*?class="browsehappy ReadPolicy">(.*?)</div>', re.I)
        content = re_policy.sub('', content)
        re_nav = re.compile(r'<div class="nav-gblnav hidden-xs  hidden-sm">(.*?)</div>', re.I)
        content = re_nav.sub('', content)
        re_affix = re.compile(r'<header class="affix-top">(.*?)</header>', re.I)
        content = re_affix.sub('', content)
        re_footer = re.compile(r'<footer>(.*?)</footer>', re.I)
        content = re_footer.sub('', content)
        re_con = re.compile(r'<div class="container">(.*?)</div>', re.I)
        content = re_con.sub('', content)
        re_global = re.compile(r'<div class="worldwide-language">(.*?)</div>', re.I)
        content = re_global.sub('', content)
        re_blank_next = re.compile(r'\s{3}', re.I)
        content = re_blank_next.sub('bupt', content)
        soup = BeautifulSoup(content, 'lxml')
        contents = soup.text
        sentences = contents.split('bupt')
        for sentence in sentences:
            if not sentence.strip() == '':
                sentence = sentence.strip(' ')
                print(sentence)
        '''
        for sentence in sentences:
            if not sentence.strip() == '':
                f = codecs.open('en_1.txt', 'a', 'utf-8')
                sentence = sentence.strip(' ')
                if sentence != '|' and sentence != '×':
                    f.write(sentence + '\n')
                    f.close()

        document = Document('test.docx')
        table = document.tables[0]
        now =1
        for sentence in sentences:
            if not sentence.strip() == '':
                rows = len(table.rows)
                sentence = sentence.strip(' ')
                if sentence != '|' and sentence != '×':
                    if now > rows:
                        table.add_row()
                    table.cell(now, 1).text = sentence
                    now = now+1
        document.save('test.docx')

        def read_txt():
            content = ''
            f = codecs.open('en.txt', 'a', 'utf-8')
            content = f.read()
            return content

        def translate(content):
            youdao_url = 'http://fanyi.youdao.com/translate?smartresult=dict&smartresult=rule'
            data = {}

            data['i'] = content
            data['from'] = 'AUTO'
            data['to'] = 'AUTO'
            data['smartresult'] = 'dict'
            data['client'] = 'fanyideskweb'
            data['salt'] = '1525141473246'
            data['sign'] = '47ee728a4465ef98ac06510bf67f3023'
            data['doctype'] = 'json'
            data['version'] = '2.1'
            data['keyfrom'] = 'fanyi.web'
            data['action'] = 'FY_BY_CLICKBUTTION'
            data['typoResult'] = 'false'
            data = urllib.parse.urlencode(data).encode('utf-8')

            youdao_response = urllib.request.urlopen(youdao_url, data)
            youdao_html = youdao_response.read().decode('utf-8')
            target = json.loads(youdao_html)

            trans = target['translateResult']
            ret = ''
            for i in range(len(trans)):
                line = ''
                for j in range(len(trans[i])):
                    line = trans[i][j]['tgt']
                ret += line + '\n'

            return ret

        def write(content):
            with open(en_tran.txt, 'a+') as f:
                f.write(content)
                '''

